from __future__ import annotations

from pathlib import Path

from ..errors import MissingDependencyError
from ..models import Document, Section, SourceInfo


class DocxAdapter:
    name = "docx"

    def supports(self, source: Path) -> bool:
        return source.suffix.lower() == ".docx"

    def load(self, source: Path) -> Document:
        try:
            from docx import Document as DocxDocument
        except ImportError as exc:  # pragma: no cover - environment dependent
            raise MissingDependencyError(
                "DOCX input requires: pip install 'ssmdconvert[docx]'"
            ) from exc
        docx = DocxDocument(str(source))
        lines: list[str] = []
        for paragraph in docx.paragraphs:
            text = paragraph.text.strip()
            if not text:
                if lines and lines[-1] != "":
                    lines.append("")
                continue
            style = (paragraph.style.name if paragraph.style else "").lower()
            if style.startswith("heading"):
                try:
                    level = int(style.rsplit(" ", 1)[-1])
                except ValueError:
                    level = 1
                lines.extend(["#" * max(1, min(level, 6)) + " " + text, ""])
            else:
                lines.extend([text, ""])
        props = docx.core_properties
        metadata = {"title": props.title or source.stem}
        if props.author:
            metadata["author"] = props.author
        return Document(
            source=SourceInfo(
                format="docx",
                media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                path=source,
                name=source.name,
            ),
            sections=[Section("section-0001", "\n".join(lines).strip())],
            metadata=metadata,
        )
